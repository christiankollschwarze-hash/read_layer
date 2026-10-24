"""READ LAYER — Kontrollansicht Woertliche Rede als Word-Datei (CK 14.10.2026)
Vor dem Hoerspiel-Overlay: Figuren mit Steckbrief + Farbe, Text mit farbig markierter woertlicher Rede.
Input : work/casting.json (Steckbriefe, Stage 57), work/sprecher.json (Stage 57b), work/400_regie.json
Output: work/Kontrollansicht_<lang>.docx
Aufruf: python3 tools/kontrollansicht_docx.py <book_id> [de|en]
Farben wie TAELLA-App: Frauen Rottoene, Maenner Blautoene (je Figur eigene Nuance), Erzaehler schwarz.
Unsichere Zuordnungen (57b "u") sind gelb hinterlegt und mit (?) markiert.
"""
import colorsys, json, os, sys
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from config import config  # noqa

ERZ = "Erzaehler"


def _hash(name):
    h = 0
    for c in name:
        h = (h * 31 + ord(c)) & 0xFFFFFFFF
    return h


def farbe(name, g):
    """gleiche Logik wie taella-app/src/components/SprecherLabel.jsx (Helligkeit 36 %)."""
    x = _hash(name)
    if g == "w":
        h, s = (345 + x % 30) % 360, 55 + (x % 4) * 8
    elif g == "m":
        h, s = 198 + x % 38, 50 + (x % 4) * 9
    else:
        h, s = 0, 0
    r, gg, b = colorsys.hls_to_rgb(h / 360, 0.36, s / 100)
    return RGBColor(int(r * 255), int(gg * 255), int(b * 255))


def _zelle_farbe(zelle, rgb):
    tcPr = zelle._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "%02X%02X%02X" % (rgb[0], rgb[1], rgb[2])); tcPr.append(shd)


def main(book_id, lang="de"):
    wd = config.work_dir(book_id)
    data = json.load(open(config.work_path(book_id, "400_regie"), encoding="utf-8"))
    cast = json.load(open(os.path.join(wd, "casting.json"), encoding="utf-8"))
    spr = json.load(open(os.path.join(wd, "sprecher.json"), encoding="utf-8"))
    try:
        reg = json.load(open(os.path.join(ROOT, "stimmen_archiv", "stimmen.json"), encoding="utf-8"))
        reg = reg.get("stimmen", reg)
    except Exception:
        reg = {}
    figuren = cast.get("steckbriefe", {}).get("figuren", [])
    plan = {e["figur"]: e for e in (cast.get("plan") or {}).get("ensemble", [])}
    geschl = {f["name"]: (f.get("merkmale") or {}).get("geschlecht", "") for f in figuren}
    # Anteil (Segmente) je Figur aus sprecher.json
    anteil = {}
    for e in spr.values():
        for x in e.get(lang, []):
            anteil[x["s"]] = anteil.get(x["s"], 0) + 1

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Georgia"; st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
    meta = data.get("meta", {})
    doc.add_heading(f"Kontrollansicht wörtliche Rede – {meta.get('title', book_id)} ({lang.upper()})", 0)
    doc.add_paragraph("Farbig = Figurenrede laut Sprecherzuordnung (Stage 57b). Schwarz = Erzähler. "
                      "Gelb hinterlegt mit (?) = unsichere Zuordnung. Korrekturen bitte direkt im Text vermerken "
                      "(z. B. Kommentar mit dem richtigen Namen).")

    doc.add_heading("Figuren", 1)
    t = doc.add_table(rows=1, cols=7); t.style = "Light Grid Accent 1"
    for i, h in enumerate(["", "Figur", "Geschlecht / Alter", "Temperament", "Dialekt / Akzent", "Besonderheiten", "Stimme · Anteil"]):
        t.rows[0].cells[i].text = h
    figs = sorted(figuren, key=lambda f: -anteil.get(f["name"], 0))
    for f in figs:
        m = f.get("merkmale") or {}
        c = farbe(f["name"], m.get("geschlecht"))
        row = t.add_row().cells
        _zelle_farbe(row[0], c)
        p = row[1].paragraphs[0]; r = p.add_run(f["name"]); r.bold = True; r.font.color.rgb = c
        if f.get("name_de") and f.get("name_de") != f["name"]:
            p.add_run(f"\n{f['name_de']}").italic = True
        row[2].text = f"{ {'m':'männlich','w':'weiblich'}.get(m.get('geschlecht'), m.get('geschlecht') or '–') } · {m.get('alter','–')}"
        row[3].text = ", ".join(m.get("temperament") or []).replace("ue", "ü").replace("oe", "ö").replace("ae", "ä") or "–"
        row[4].text = m.get("akzent") or "–"
        row[5].text = (f.get("kurz") or "") + (f"\nStimme: {f.get('stimm_beschreibung_en')}" if f.get("stimm_beschreibung_en") else "")
        e = plan.get(f["name"], {}); sn = e.get("stimme")
        ruf = reg.get(sn, {}).get("rufname") if sn else None
        row[6].text = (f"{ruf} ({sn})" if ruf else (sn or "Erzähler-Stimme")) + f"\n{anteil.get(f['name'], 0)} Redeteile"
    for row in t.rows:
        row.cells[0].width = Cm(0.6)

    doc.add_heading("Text", 1)
    unsicher = 0
    absatz = None
    for tid, e in data["content"].items():
        typ = e.get("type")
        if typ in ("skip", "tt"):
            continue
        if typ == "p":
            absatz = None
            continue
        text = (e.get(lang) or "").strip()
        if not text:
            continue
        if typ == "h":
            doc.add_heading(text, 2); absatz = None
            continue
        if absatz is None:
            absatz = doc.add_paragraph()
        segs = (spr.get(tid) or {}).get(lang) or [{"t": text, "s": ERZ}]
        for x in segs:
            r = absatz.add_run(x["t"])
            if x["s"] != ERZ:
                r.font.color.rgb = farbe(x["s"], geschl.get(x["s"]))
            if x.get("u"):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                q = absatz.add_run(f" (? {x['s']})"); q.font.size = Pt(8); q.font.highlight_color = WD_COLOR_INDEX.YELLOW
                unsicher += 1
            if x is not segs[-1]:
                absatz.add_run(" ")
        absatz.add_run(" ")
    out = os.path.join(wd, f"Kontrollansicht_{lang}.docx")
    doc.save(out)
    print(f"[kontrolle] {len(figuren)} Figuren, {unsicher} unsichere Stellen -> {out}")


if __name__ == "__main__":
    a = [x for x in sys.argv[1:] if not x.startswith("--")]
    main(a[0], a[1] if len(a) > 1 else "de")
